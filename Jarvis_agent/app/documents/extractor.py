import io
import csv
from typing import List, Dict, Any, Tuple
from pathlib import Path
from pypdf import PdfReader
import docx
from app.core.exceptions import DocumentProcessingError
from app.core.config import settings


class DocumentExtractor:
    """Extracts text content and page metadata from PDF, DOCX, TXT, and CSV files."""

    ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".csv"}

    @classmethod
    def validate_file(cls, filename: str, content_bytes: bytes) -> None:
        """Validate file extension and size."""
        ext = Path(filename).suffix.lower()
        if ext not in cls.ALLOWED_EXTENSIONS:
            raise DocumentProcessingError(
                f"Unsupported file type '{ext}'. Supported formats: {', '.join(cls.ALLOWED_EXTENSIONS)}"
            )

        max_bytes = settings.MAX_FILE_SIZE_MB * 1024 * 1024
        if len(content_bytes) > max_bytes:
            raise DocumentProcessingError(
                f"File size ({len(content_bytes) / 1024 / 1024:.2f} MB) exceeds maximum allowed limit ({settings.MAX_FILE_SIZE_MB} MB)."
            )

        if len(content_bytes) == 0:
            raise DocumentProcessingError("File is empty.")

    @classmethod
    def extract_text(cls, filename: str, content_bytes: bytes) -> Tuple[str, List[Dict[str, Any]]]:
        """
        Extract text from file.
        Returns:
            Tuple of (full_normalized_text, list of page/chunk sections: [{"page": int, "text": str}])
        """
        cls.validate_file(filename, content_bytes)
        ext = Path(filename).suffix.lower()

        try:
            if ext == ".pdf":
                return cls._extract_pdf(content_bytes)
            elif ext == ".docx":
                return cls._extract_docx(content_bytes)
            elif ext == ".csv":
                return cls._extract_csv(content_bytes)
            elif ext == ".txt":
                return cls._extract_txt(content_bytes)
            else:
                raise DocumentProcessingError(f"Unsupported format: {ext}")
        except DocumentProcessingError:
            raise
        except Exception as e:
            raise DocumentProcessingError(f"Failed to extract document contents: {str(e)}")

    @classmethod
    def _extract_pdf(cls, content_bytes: bytes) -> Tuple[str, List[Dict[str, Any]]]:
        reader = PdfReader(io.BytesIO(content_bytes))
        sections = []
        full_text_parts = []

        total_pages = len(reader.pages)
        if total_pages == 0:
            raise DocumentProcessingError("PDF contains no pages.")

        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            page_text = page_text.strip()
            if not page_text:
                page_text = "[Notice: Scanned page or empty text layer. OCR not configured.]"
            sections.append({"page": idx + 1, "text": page_text})
            full_text_parts.append(f"--- Page {idx + 1} ---\n{page_text}")

        full_text = "\n\n".join(full_text_parts)
        if len(full_text) > settings.MAX_EXTRACT_CHARS:
            full_text = full_text[:settings.MAX_EXTRACT_CHARS] + "\n[Content truncated at limit]"

        return full_text, sections

    @classmethod
    def _extract_docx(cls, content_bytes: bytes) -> Tuple[str, List[Dict[str, Any]]]:
        doc = docx.Document(io.BytesIO(content_bytes))
        paragraphs = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                paragraphs.append(text)

        # Also extract table contents
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n\n".join(paragraphs)
        if not full_text.strip():
            full_text = "[Empty document]"

        if len(full_text) > settings.MAX_EXTRACT_CHARS:
            full_text = full_text[:settings.MAX_EXTRACT_CHARS] + "\n[Content truncated at limit]"

        sections = [{"page": 1, "text": full_text}]
        return full_text, sections

    @classmethod
    def _extract_txt(cls, content_bytes: bytes) -> Tuple[str, List[Dict[str, Any]]]:
        try:
            text = content_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = content_bytes.decode("latin-1", errors="replace")

        text = text.strip()
        if not text:
            text = "[Empty text file]"

        if len(text) > settings.MAX_EXTRACT_CHARS:
            text = text[:settings.MAX_EXTRACT_CHARS] + "\n[Content truncated at limit]"

        return text, [{"page": 1, "text": text}]

    @classmethod
    def _extract_csv(cls, content_bytes: bytes) -> Tuple[str, List[Dict[str, Any]]]:
        try:
            raw = content_bytes.decode("utf-8")
        except UnicodeDecodeError:
            raw = content_bytes.decode("latin-1", errors="replace")

        reader = csv.reader(io.StringIO(raw))
        rows = list(reader)
        if not rows:
            return "[Empty CSV]", [{"page": 1, "text": "[Empty CSV]"}]

        formatted_lines = []
        for row_idx, row in enumerate(rows[:500]):  # cap rows
            formatted_lines.append(" | ".join(row))

        full_text = "\n".join(formatted_lines)
        if len(full_text) > settings.MAX_EXTRACT_CHARS:
            full_text = full_text[:settings.MAX_EXTRACT_CHARS] + "\n[Content truncated at limit]"

        return full_text, [{"page": 1, "text": full_text}]
